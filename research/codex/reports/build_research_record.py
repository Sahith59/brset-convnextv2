"""Build the maintained DOCX research record and concise PPTX update."""
import hashlib
import json
import zipfile
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor as DocColor
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches as PInches, Pt as PPt

ROOT = Path(__file__).resolve().parents[3]
HERE = Path(__file__).resolve().parent
DOCX = HERE / "BRSET_mBRSET_Research_Record.docx"
PPTX = HERE / "BRSET_mBRSET_Research_Update.pptx"
CHECKS = HERE / "record_checks.json"
FONT = "Times New Roman"
UPDATED = date.today().strftime("%d %B %Y").lstrip("0")

RESULTS = ROOT / "research/codex/step2/full_pool/baseline_assessment_three_seeds.json"
SEED2 = ROOT / "research/codex/step2/full_pool/baseline_assessment_seed2.json"
DESIGN = ROOT / "research/codex/step3/design_counts.json"
PREFLIGHT = ROOT / "research/codex/step3/preflight.json"
STEP3_RUNS = ROOT / "research/codex/step3/runs"
STEP2_RUNS = ROOT / "research/codex/step2/full_pool/runs"
STEP3_LAUNCH = ROOT / "research/codex/step3/launch_seed0.json"
STEP3_SEED0 = ROOT / "research/codex/step3/seed0_validation_summary.json"
STEP3_C1_CROSS = ROOT / "research/codex/step3/c1_validation_three_seeds.json"
STEP3_FINAL = ROOT / "research/codex/step3/final_curve_review.json"
STEP4_AUDIT = ROOT / "research/codex/step4/appearance_label_audit.json"
STEP4_REFIT = ROOT / "research/codex/step4/degradation_refit.json"
STEP4_PREFLIGHT = ROOT / "research/codex/step4/preflight.json"
STEP4_SEED0 = ROOT / "research/codex/step4/seed0_validation_summary.json"


def sha(path):
    h = hashlib.sha256()
    with open(path, "rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            h.update(block)
    return h.hexdigest()


def set_doc_run(run, size=10, bold=False, italic=False, color="000000"):
    run.font.name = FONT
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), FONT)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = DocColor.from_string(color)


def doc_text(paragraph, text, size=10, bold=False, italic=False):
    run = paragraph.add_run(text)
    set_doc_run(run, size, bold, italic)
    return run


def style_doc_paragraph(paragraph, before=0, after=4, line=1.0):
    paragraph.paragraph_format.space_before = Pt(before)
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.line_spacing = line


def add_heading(document, number, title, level=1):
    p = document.add_paragraph()
    size = 11 if level == 1 else 10
    doc_text(p, f"{number}  {title}", size=size, bold=True, italic=(level == 2))
    style_doc_paragraph(p, before=8 if level == 1 else 5, after=3)
    return p


def add_body(document, text, bold_prefix=None):
    p = document.add_paragraph()
    if bold_prefix and text.startswith(bold_prefix):
        doc_text(p, bold_prefix, bold=True)
        doc_text(p, text[len(bold_prefix):])
    else:
        doc_text(p, text)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    style_doc_paragraph(p)
    return p


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def add_doc_table(document, rows, widths=None):
    table = document.add_table(rows=len(rows), cols=len(rows[0]))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for i, row in enumerate(rows):
        for j, value in enumerate(row):
            cell = table.cell(i, j)
            cell.text = ""
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if j else WD_ALIGN_PARAGRAPH.LEFT
            doc_text(p, str(value), size=8.5, bold=(i == 0))
            style_doc_paragraph(p, after=0)
            if i == 0:
                shade(cell, "D9E2F3")
            elif i % 2 == 0:
                shade(cell, "F3F5F7")
            if widths:
                cell.width = Inches(widths[j])
    document.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def add_page_number(section):
    p = section.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    fld_begin = OxmlElement("w:fldChar"); fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText"); instr.set(qn("xml:space"), "preserve"); instr.text = " PAGE "
    fld_end = OxmlElement("w:fldChar"); fld_end.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_begin, instr, fld_end])
    set_doc_run(run, 8)


def step3_validation_rows():
    if STEP3_SEED0.exists():
        verified = json.loads(STEP3_SEED0.read_text())
        rows = []
        for name, metrics in [("B1 existing", verified["reference"])] + [
                (arm.split("_")[0], record["metrics"]) for arm, record in verified["controls"].items()]:
            dr, me = metrics["diabetic_retinopathy"], metrics["macular_edema"]
            rows.append([name, f"{dr['f1_positive']:.4f}", f"{dr['auroc']:.4f}",
                         f"{me['f1_positive']:.4f}", f"{me['auroc']:.4f}"])
        return rows
    locations = [("B1 existing", STEP2_RUNS / "B1_seed0/selection_summary.json")]
    locations.extend((name.split("_")[0], STEP3_RUNS / f"{name}_seed0/selection_summary.json") for name in (
        "C1_equal_domain", "C2_target_label", "C3_equal_domain_target_label"))
    rows = []
    for name, path in locations:
        if not path.exists():
            continue
        data = json.loads(path.read_text()); metrics = data["metrics"]
        dr, me = metrics["diabetic_retinopathy"], metrics["macular_edema"]
        rows.append([name, f"{dr['f1_positive']:.4f}", f"{dr['auroc']:.4f}",
                     f"{me['f1_positive']:.4f}", f"{me['auroc']:.4f}"])
    return rows


def step3_state(preflight):
    controls = step3_validation_rows()
    completed = {row[0] for row in controls}
    if STEP3_FINAL.exists() and json.loads(STEP3_FINAL.read_text())["decision"]["step3_complete"]:
        return "Step 3 complete; natural joint B1 retained for Step 4", controls
    if STEP3_C1_CROSS.exists():
        return "C1 replication verified across three seeds; review Step-3 closure", controls
    if {"C1", "C2", "C3"}.issubset(completed):
        if (ROOT / "research/codex/step3/seed0_validation_summary.json").exists():
            return "seed-0 controls verified; review the replication screen", controls
        return "seed-0 controls complete; validation results require verification", controls
    if STEP3_LAUNCH.exists():
        return "seed-0 controls submitted/running; no control result claimed yet", controls
    return ("preflight passed; launch pending" if preflight and preflight.get("pass") else "preflight pending"), controls


def build_docx(summary, design, preflight):
    run_state, control_rows = step3_state(preflight)
    document = Document()
    section = document.sections[0]
    section.top_margin = section.bottom_margin = Inches(0.70)
    section.left_margin = section.right_margin = Inches(0.75)
    for name, size, bold, italic in [("Normal", 10, False, False), ("Title", 18, True, False)]:
        style = document.styles[name]
        style.font.name = FONT; style.font.size = Pt(size); style.font.bold = bold; style.font.italic = italic
        style.element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), FONT)
    add_page_number(section)

    p = document.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc_text(p, "BRSET to mBRSET Cross-Device Retinal Classification", 18, True)
    style_doc_paragraph(p, after=2)
    p = document.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc_text(p, "Living Research Record and Evidence Audit", 11, True)
    style_doc_paragraph(p, after=2)
    p = document.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if STEP4_SEED0.exists(): step4_state = "Step 4 seed-0 validation screen complete"
    elif STEP4_PREFLIGHT.exists(): step4_state = "Step 4 transform refit and preflight complete"
    elif STEP4_REFIT.exists(): step4_state = "Step 4 transform refit complete"
    else: step4_state = "Step 4 transform refit active"
    doc_text(p, f"Updated {UPDATED} | Current state: {step4_state}", 9, italic=True)
    style_doc_paragraph(p, after=8)

    p = document.add_paragraph(); doc_text(p, "Abstract—", 10, True, True)
    doc_text(p, "This record tracks a supervised cross-device study for image-level diabetic retinopathy (DR) and macular edema (ME) classification. The evidence establishes a strong joint BRSET–mBRSET baseline and closes the balance/exposure investigation. The Step-4 seed-0 screen found that target-calibrated source transformations did not pass the frozen replication gate. A novel method has not yet been established. Numerical claims are linked to audited aggregate artifacts, and known limitations are retained rather than removed from later updates.")
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY; style_doc_paragraph(p, after=7)

    add_heading(document, "I.", "RESEARCH QUESTION")
    add_body(document, "Can a model trained with tabletop-camera BRSET images and labeled handheld-camera mBRSET images improve DR and ME recognition on mBRSET, and can the improvement be attributed to a defensible cross-device mechanism rather than data quantity, class balance, threshold choice, or random seed?")
    add_body(document, "Current contribution status: The study has a verified experimental foundation, a repeated joint-training advantage, and a closed balance/exposure analysis. A paper-level novelty claim remains conditional on the preservation-constrained appearance experiments and closest-method evidence.", "Current contribution status:")

    add_heading(document, "II.", "DATA AND EVALUATION BOUNDARY")
    add_doc_table(document, [
        ["Dataset role", "Images", "Patients", "Use"],
        ["BRSET train", "11,372", "5,966", "Source fitting only"],
        ["mBRSET train", "3,402", "899", "Target fitting"],
        ["mBRSET validation", "725", "193", "Checkpoint and threshold selection"],
        ["mBRSET test", "732", "193", "Gated assessment; historically reused"],
    ], [2.0, 1.0, 1.0, 3.0])
    add_body(document, "Labels are image-level binary outcomes: any DR and ME. A patient is counted as positive in prevalence summaries when at least one image is positive; this is not a separately adjudicated patient diagnosis. Exact native-RGB and EXIF-normalized duplicate scans passed for the included images, but near-duplicate and cross-dataset identity risks are not fully excluded.")

    add_heading(document, "III.", "DEGRADATION STUDY COMPLETED FOR ADVISOR UPDATE")
    add_body(document, "The fitted transform was estimated on training-only images and then applied, without refitting, to a patient-disjoint training-only check set. It changes image appearance; it has not yet been shown to preserve lesions or improve diagnosis.")
    p = document.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc_text(p, "D(x,t) = sqrt[Σₖ ((sₖ(x) − μₖ,t) / σₖ,t)²]", 10, italic=True)
    style_doc_paragraph(p, after=4)
    add_body(document, "Here sₖ(x) is one measured appearance statistic for image x, while μₖ,t and σₖ,t are the target-domain mean and scale. Lower distance means closer agreement with the selected appearance statistics; it does not mean greater clinical realism.")
    add_doc_table(document, [
        ["Preprocessing", "BRSET unchanged", "Local generic transform", "Frozen fitted transform"],
        ["Legacy 512 resize", "6.742", "10.092", "1.261"],
        ["560 resize / 512 crop", "8.474", "12.011", "4.294"],
    ])
    add_body(document, "Result: the frozen fitted transform reduced the measured appearance distance in both preprocessing settings. Limitation: the target examples are unpaired, and the measured statistics do not validate diagnostic preservation or classifier improvement.", "Result:")

    add_heading(document, "IV.", "FIXED BASELINE PROTOCOL")
    add_body(document, "All Step-2 arms use ConvNeXt V2 Tiny, 5,775 optimizer updates, effective batch size 64, focal loss, ordinary spatial/color augmentation, Mixup, AdamW, exponential moving average weights, and four-flip evaluation. The mBRSET validation split selects the checkpoint by mean DR/ME AUROC and selects one F1 threshold per label. No degradation, routing, or private encoder is used.")
    add_doc_table(document, [
        ["Arm", "Training data", "Meaning"],
        ["B0 target-only", "3,402 mBRSET", "Learn only from labeled target images"],
        ["B1 joint", "11,372 BRSET + 3,402 mBRSET", "Mix source and target in one training pool"],
        ["B2 source to target", "11,372 BRSET, then 3,402 mBRSET", "Pretrain on source, then fine-tune on target"],
    ])

    add_heading(document, "V.", "STEP-2 RESULTS: THREE TRAINING SEEDS")
    rows = [["Seed", "Arm", "DR F1", "DR AUROC", "ME F1", "ME AUROC"]]
    for seed_index, seed in enumerate(summary["seeds"]):
        for arm in ("B0", "B1", "B2"):
            aggregate = summary["aggregate"][arm]
            dr = {metric: aggregate["diabetic_retinopathy"][metric]["values"][seed_index]
                  for metric in ("f1_positive", "auroc")}
            me = {metric: aggregate["macular_edema"][metric]["values"][seed_index]
                  for metric in ("f1_positive", "auroc")}
            rows.append([seed, arm, f"{dr['f1_positive']:.4f}", f"{dr['auroc']:.4f}",
                         f"{me['f1_positive']:.4f}", f"{me['auroc']:.4f}"])
    add_doc_table(document, rows)
    means = summary["aggregate"]
    rows = [["Arm", "DR F1 mean ± SD", "DR AUROC", "ME F1 mean ± SD", "ME AUROC"]]
    for arm in ("B0", "B1", "B2"):
        def fmt(label, metric):
            x = means[arm][label][metric]
            return f"{x['mean']:.4f} ± {x['sample_sd']:.4f}"
        rows.append([arm, fmt("diabetic_retinopathy", "f1_positive"), fmt("diabetic_retinopathy", "auroc"),
                     fmt("macular_edema", "f1_positive"), fmt("macular_edema", "auroc")])
    add_doc_table(document, rows)
    add_body(document, "Verified finding: B1 exceeded B0 F1 for DR and ME in every seed. Mean B1−B0 differences were +0.0233 for DR and +0.0379 for ME. B1 passes the prespecified engineering screen and becomes the Step-3 reference.", "Verified finding:")
    add_body(document, "Interpretation limit: Three seeds provide a limited estimate of training variability. Patient bootstrap intervals for a fixed trained model do not include all training uncertainty. The reused test split is development benchmark evidence, not untouched external validation. The result is neither a novelty claim nor a clinical superiority claim.", "Interpretation limit:")

    add_heading(document, "VI.", "STEP-3 BALANCE AND EXPOSURE CONTROLS")
    add_body(document, "Purpose: determine whether B1 benefits from domain mixing, the DR/ME label composition, or their interaction. Every model and optimization setting remains fixed; only the deterministic sampler changes.")
    add_doc_table(document, [
        ["Arm", "BRSET/mBRSET sampling", "Joint DR/ME sampling", "Question"],
        ["B1 existing", "Natural pool ratio", "Natural", "Reference"],
        ["C1", "50% / 50%", "Natural per domain", "Effect of domain frequency"],
        ["C2", "Natural pool ratio", "Target proportions per domain", "Effect of label composition"],
        ["C3", "50% / 50%", "Target proportions per domain", "Interaction"],
    ])
    add_body(document, "Exposure fact: over 369,600 image draws, B0 sees about 369,600 target draws, natural B1 about 85,108, and an equal-domain control 184,800. B1 therefore did not improve simply because it saw more target draws than B0. This does not reveal why source data helped.")
    add_body(document, "Small-cell risk: DR0/ME1 contains only 22 BRSET and 14 mBRSET training images. Repetition is reported explicitly and cannot be treated as new independent evidence. A sensitivity analysis is required before adopting a target-label-matched sampler as the paper baseline.")
    state = "passed" if preflight and preflight.get("pass") else "pending"
    add_body(document, f"Implementation state: deterministic exact-quota sampler, resume equivalence tests and exposure ledger prepared. Allocated preflight: {state}; current run state: {run_state}. No Step-3 test assessment has been performed.", "Implementation state:")
    if len(control_rows) > 1:
        add_doc_table(document, [["Arm", "Validation DR F1", "DR AUROC", "ME F1", "ME AUROC"]] + control_rows)
        add_body(document, "These are validation-screening values. They must be verified and cannot be reported as test results.")
    if STEP3_C1_CROSS.exists():
        cross = json.loads(STEP3_C1_CROSS.read_text())
        cross_rows = [["Seed", "B1 DR F1", "C1 DR F1", "Difference", "B1 ME F1", "C1 ME F1", "Difference"]]
        for seed in cross["seeds"]:
            r = cross["per_seed"][str(seed)]
            cross_rows.append([seed, f"{r['B1']['diabetic_retinopathy']['f1_positive']:.4f}",
                f"{r['C1']['diabetic_retinopathy']['f1_positive']:.4f}", f"{r['C1-B1']['diabetic_retinopathy']['f1_positive']:+.4f}",
                f"{r['B1']['macular_edema']['f1_positive']:.4f}", f"{r['C1']['macular_edema']['f1_positive']:.4f}",
                f"{r['C1-B1']['macular_edema']['f1_positive']:+.4f}"])
        add_doc_table(document, cross_rows)
        add_body(document, "C1 cross-seed values are validation-only evidence used to choose the Step-4 baseline. They are not an external test result.")
    if STEP3_FINAL.exists():
        final = json.loads(STEP3_FINAL.read_text())
        dr = final["c1_minus_b1_validation"]["diabetic_retinopathy"]
        me = final["c1_minus_b1_validation"]["macular_edema"]
        add_body(document, f"Step-3 decision: retain natural B1. Mean C1−B1 validation differences were {dr['f1_difference_mean']:+.4f} DR F1 and {me['f1_difference_mean']:+.4f} ME F1. DR changed direction across seeds, while C1 mean AUROC was lower for both labels. All six B1/C1 runs peaked before the final quarter and declined thereafter, so the conditional longer-exposure run is not justified.", "Step-3 decision:")

    add_heading(document, "VII.", "STEP-4 TRAINING-ONLY APPEARANCE AUDIT")
    if STEP4_AUDIT.exists():
        audit = json.loads(STEP4_AUDIT.read_text())
        audit_rows = [["Statistic", "Raw gap / target SD", "Common-composition gap / target SD", "Change"]]
        for name in ("sharpness", "brightness", "contrast", "falloff", "saturation"):
            entry = audit["gap_sensitivity"][name]
            raw = entry["raw_gap_in_target_sd"]
            standardized = entry["standardized_gap_in_target_sd"]
            audit_rows.append([name.capitalize(), f"{raw:+.4f}", f"{standardized:+.4f}", f"{standardized-raw:+.4f}"])
        add_body(document, "Purpose: check whether unequal DR/ME prevalence materially changes the five measured BRSET–mBRSET appearance gaps before refitting the degradation transform.")
        add_doc_table(document, audit_rows)
        add_body(document, "Verified finding: all 14,774 original-training images were measured. Standardizing both domains to the same pooled joint-label composition reversed no direction and changed each gap by less than 0.05 mBRSET standard deviations. Proceed with a global training-only refit; a label-standardized classifier arm is not justified.", "Verified finding:")
        add_body(document, "Interpretation limit: this is descriptive sensitivity analysis. Recorded DR/ME labels do not separate camera effects from population, site, image quality, other disease or unmeasured differences, and global appearance summaries do not establish lesion preservation.", "Interpretation limit:")
    else:
        add_body(document, "The training-only appearance/label-composition audit is pending. No Step-4 classifier or test assessment has started.")

    if STEP4_REFIT.exists():
        refit = json.loads(STEP4_REFIT.read_text())
        add_heading(document, "VII-A.", "STEP-4 TRANSFORM REFIT", level=2)
        rows = [["Validation transform", "Appearance distance mean", "Range across 3 draws"]]
        for label, record in [
            ("Unchanged source", refit["identity_validation"]),
            ("Historical frozen", refit["historical_frozen_validation"]),
            ("Refitted full", refit["arms"]["full"]["validation"]),
            ("Refitted no overlays", refit["arms"]["overlay_free"]["validation"]),
        ]:
            rows.append([label, f"{record['distance_mean']:.3f}", f"{record['distance_min']:.3f}–{record['distance_max']:.3f}"])
        add_doc_table(document, rows)
        add_body(document, "The refit used target means/scales from all 3,402 target-training images and two patient-disjoint 96-image source subsets for fitting and validation. These distances measure five global statistics; they do not certify lesions or camera physics.")
        full = refit["arms"]["full"]["params"]
        clean = refit["arms"]["overlay_free"]["params"]
        add_doc_table(document, [
            ["Parameter", "Refitted full", "Refitted no overlays"],
            ["Blur sigma", f"{full['blur_sigma']:.4f}", f"{clean['blur_sigma']:.4f}"],
            ["Illumination strength", f"{full['light_strength']:.4f}", f"{clean['light_strength']:.4f}"],
            ["Spot / hole count; halo", f"{int(full['n_spot'])} / {int(full['n_hole'])}; {full['halo']:.4f}", "0 / 0; 0"],
            ["Brightness / contrast / saturation", f"{full['brightness_gain']:.4f} / {full['contrast_gain']:.4f} / {full['saturation_gain']:.4f}", f"{clean['brightness_gain']:.4f} / {clean['contrast_gain']:.4f} / {clean['saturation_gain']:.4f}"],
            ["Sensor-noise sigma", f"{full['noise_sigma']:.5f}", f"{clean['noise_sigma']:.5f}"],
        ])
        if STEP4_PREFLIGHT.exists():
            gate = json.loads(STEP4_PREFLIGHT.read_text())
            add_body(document, f"Implementation gate: {'passed' if gate['pass'] else 'failed'}. The target-domain training transform is bitwise identical to B1, fitted parameters are valid, and a private four-row qualitative gallery was generated. The gallery is excluded from this report under the privacy policy.", "Implementation gate:")
    if STEP4_SEED0.exists():
        screen = json.loads(STEP4_SEED0.read_text())
        rows = [["Arm", "DR F1", "Δ from B1", "DR AUROC", "ME F1", "Δ from B1", "ME AUROC", "Replicate?"]]
        ref = screen["reference"]["metrics"]
        rows.append(["B1", f"{ref['diabetic_retinopathy']['f1_positive']:.4f}", "—", f"{ref['diabetic_retinopathy']['auroc']:.4f}", f"{ref['macular_edema']['f1_positive']:.4f}", "—", f"{ref['macular_edema']['auroc']:.4f}", "Reference"])
        for arm, record in screen["arms"].items():
            metric, delta = record["metrics"], record["delta_from_B1"]
            rows.append([arm, f"{metric['diabetic_retinopathy']['f1_positive']:.4f}", f"{delta['diabetic_retinopathy']['f1_positive']:+.4f}", f"{metric['diabetic_retinopathy']['auroc']:.4f}", f"{metric['macular_edema']['f1_positive']:.4f}", f"{delta['macular_edema']['f1_positive']:+.4f}", f"{metric['macular_edema']['auroc']:.4f}", "Yes" if record["replication_screen_pass"] else "No"])
        add_doc_table(document, rows)
        add_body(document, "Step-4 seed-0 decision: no arm passed the frozen replication gate. Each DR F1 change was positive but below +0.01, while ME was essentially flat or slightly lower. A3 improved DR AUROC by +0.0038, but this single-seed signal is insufficient for a gain claim. No Step-4 test assessment is included.", "Step-4 seed-0 decision:")

    add_heading(document, "VIII.", "NOVELTY GATE, WORLD MODEL AND PAPER DIRECTION")
    add_body(document, "Candidate method: a lightweight target-appearance augmentation fitted on training data, combined with an explicit diagnostic-damage or lesion-preservation constraint. Appearance matching alone is insufficient because transformations can improve global statistics while obscuring lesions.")
    add_body(document, "The novelty claim becomes supportable only if the method is distinct from the closest augmentation, consistency, synthesis and structural-preservation methods; improves over the strongest fair baseline across seeds; survives an ablation that isolates the preservation constraint; and passes a qualified preservation review or a validated lesion-sensitive proxy.")
    add_body(document, "Falsification rule: if balance controls explain the apparent gain, or fitted degradation fails to improve validation performance consistently, the mechanism will not be presented as an effective method. The paper direction must then shift to the strongest supported diagnostic finding rather than inventing a positive result.")
    add_body(document, "World-model decision: CheXWorld (CVPR 2025) is the most plausible reference behind the meeting discussion. It predicts target latent features from transformed context features conditioned on known blur/color parameters. Its official example uses ViT-Base for 300 epochs on eight RTX 4090 GPUs. A smaller fundus latent-transition objective is a conditional Step-5 candidate; full CheXWorld or GenDeg training is not justified before Step 4 validates the acquisition-variation premise.", "World-model decision:")
    add_body(document, "Advisor coverage: same-configuration comparisons, class prevalence, sampling controls, qualitative degradation examples and the seed-0 classifier screen are addressed. The fitted transform did not pass the replication gate. Diagnostic-preservation validation and a fair longer-budget routing or latent-transition study remain open. The transcript also contains a user-owned request to disclose this external collaboration to the new employer; completion has not been documented.", "Advisor coverage:")

    add_heading(document, "IX.", "NEXT ACTIONS")
    for item in [
        "Preserve A1/A2/A3 as a completed negative validation screen; do not assess them on test or replicate them automatically.",
        "Freeze one new content-preserving mechanism and a matched control before further GPU work.",
        "Choose between a bounded CheXWorld-style latent-transition objective and an explicitly lesion-aware preservation constraint using feasibility and closest-work evidence.",
        "Require repeated diagnostic benefit and an isolating ablation before writing a contribution claim.",
        "Ask Dong to confirm the exact world-model paper title and venue when he responds.",
    ]:
        p = document.add_paragraph(style=None); p.style = document.styles["Normal"]
        p.paragraph_format.left_indent = Inches(0.20); p.paragraph_format.first_line_indent = Inches(-0.15)
        doc_text(p, "•  " + item); style_doc_paragraph(p, after=2)

    add_heading(document, "X.", "AUDIT TRAIL")
    add_doc_table(document, [
        ["Date", "Milestone", "Evidence"],
        ["10 Sep 2026", "Degradation update completed and user sent it to Dong", "validation_summary.json; deliverable_checks.json"],
        ["10 Sep 2026", "Comparison protocol frozen", "STEP_1_PROTOCOL.md; preflight_v1.json"],
        ["12 Sep 2026", "Nine Step-2 runs and three assessments complete", "baseline_assessment_three_seeds.json"],
        ["12 Sep 2026", "Independent seed checks and Step-2 closure", "seed*_independent_verification.json; STEP_2_FINAL_REVIEW.md"],
        ["12 Sep 2026", "Step-3 sampler controls prepared", "SAMPLER_PROTOCOL.md; preflight.json"],
        ["13 Sep 2026", "Step-3 controls, C1 replication and curve decision complete", "c1_validation_three_seeds.json; final_curve_review.json"],
        ["14 Sep 2026", "Step-4 training-only label-composition audit complete", "appearance_label_audit.json; APPEARANCE_LABEL_AUDIT_REVIEW.md"],
        ["14 Sep 2026", "Meeting/world-model direction audited", "DONG_TRANSCRIPT_AUDIT_2026-09.md; STEP4_WORLD_MODEL_DIRECTION_REVIEW.md"],
        ["15 Sep 2026", "Step-4 seed-0 augmentation screen complete; no arm passed replication gate", "seed0_validation_summary.json; SEED0_RESULTS_REVIEW.md"],
    ])
    document.save(DOCX)


BLUE = RGBColor(31, 78, 121); INK = RGBColor(30, 30, 30); MUTED = RGBColor(90, 90, 90)
PALE = RGBColor(235, 241, 247); GREEN = RGBColor(36, 105, 67); LIGHT_GREEN = RGBColor(226, 239, 229)
AMBER = RGBColor(154, 103, 16); LIGHT_AMBER = RGBColor(250, 241, 222); WHITE = RGBColor(255, 255, 255)


def ppt_run(run, size=18, bold=False, color=INK, italic=False):
    run.font.name = FONT; run.font.size = PPt(size); run.font.bold = bold; run.font.italic = italic
    run.font.color.rgb = color


def textbox(slide, x, y, w, h, text, size=18, bold=False, color=INK, align=PP_ALIGN.LEFT):
    shape = slide.shapes.add_textbox(PInches(x), PInches(y), PInches(w), PInches(h))
    frame = shape.text_frame; frame.clear(); frame.word_wrap = True
    p = frame.paragraphs[0]; p.alignment = align
    run = p.add_run(); run.text = text; ppt_run(run, size, bold, color)
    return shape


def ppt_box(slide, x, y, w, h, text, fill, line=BLUE, size=17, bold=False, color=INK):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, PInches(x), PInches(y), PInches(w), PInches(h))
    shape.fill.solid(); shape.fill.fore_color.rgb = fill; shape.line.color.rgb = line
    frame = shape.text_frame; frame.clear(); frame.word_wrap = True; frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = frame.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    run = p.add_run(); run.text = text; ppt_run(run, size, bold, color)
    return shape


def slide_title(slide, title, subtitle=None):
    textbox(slide, .55, .28, 12.2, .5, title, 27, True, BLUE)
    if subtitle: textbox(slide, .58, .82, 12.0, .35, subtitle, 12, False, MUTED)
    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, PInches(.58), PInches(1.23), PInches(12.1), PInches(.02))
    line.fill.solid(); line.fill.fore_color.rgb = BLUE; line.line.fill.background()


def bullets(slide, items, x=.75, y=1.55, w=11.8, h=5.2, size=18):
    h = min(h, 6.95 - y)
    shape = slide.shapes.add_textbox(PInches(x), PInches(y), PInches(w), PInches(h))
    frame = shape.text_frame; frame.clear(); frame.word_wrap = True
    for i, item in enumerate(items):
        p = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
        p.text = item; p.level = 0; p.space_after = PPt(12)
        for run in p.runs: ppt_run(run, size)
    return shape


def ppt_table(slide, rows, x, y, w, h, widths=None, size=13):
    shape = slide.shapes.add_table(len(rows), len(rows[0]), PInches(x), PInches(y), PInches(w), PInches(h))
    table = shape.table
    if widths:
        total = sum(widths)
        for j, width in enumerate(widths): table.columns[j].width = PInches(w * width / total)
    for i, row in enumerate(rows):
        for j, value in enumerate(row):
            cell = table.cell(i, j); cell.text = str(value); cell.margin_left = cell.margin_right = PInches(.06)
            cell.margin_top = cell.margin_bottom = PInches(.03); cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.fill.solid(); cell.fill.fore_color.rgb = BLUE if i == 0 else WHITE if i % 2 else PALE
            for p in cell.text_frame.paragraphs:
                p.alignment = PP_ALIGN.CENTER if j else PP_ALIGN.LEFT
                for run in p.runs: ppt_run(run, size, i == 0, WHITE if i == 0 else INK)
    return shape


def build_pptx(summary, preflight):
    run_state, control_rows = step3_state(preflight)
    prs = Presentation(); prs.slide_width = PInches(13.333); prs.slide_height = PInches(7.5)
    blank = prs.slide_layouts[6]

    s = prs.slides.add_slide(blank)
    textbox(s, .75, 1.25, 11.8, .8, "BRSET to mBRSET Cross-Device Classification", 30, True, BLUE, PP_ALIGN.CENTER)
    state = "Step 4 seed-0 screen complete" if STEP4_SEED0.exists() else "Step 4 transform refit/preflight active"
    textbox(s, 1.2, 2.25, 10.9, .6, f"Evidence update: {state}", 20, False, INK, PP_ALIGN.CENTER)
    ppt_box(s, 2.15, 3.35, 9.0, 1.05, "Natural joint training remains the strongest simple reference.\nA novel method has not yet been established.", LIGHT_GREEN, GREEN, 20, True)
    textbox(s, 1.0, 6.55, 11.3, .35, UPDATED, 12, False, MUTED, PP_ALIGN.CENTER)

    s = prs.slides.add_slide(blank); slide_title(s, "Research question and evidence chain")
    labels = [("Camera and population shift", .55), ("Strong ordinary baselines", 3.75),
              ("Balance and exposure controls", 6.95), ("Preservation-constrained method", 10.15)]
    for text, x in labels: ppt_box(s, x, 2.15, 2.6, 1.15, text, PALE, BLUE, 16, True)
    for x in (3.25, 6.45, 9.65):
        arrow = s.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, PInches(x), PInches(2.48), PInches(.42), PInches(.45))
        arrow.fill.solid(); arrow.fill.fore_color.rgb = BLUE; arrow.line.fill.background()
    bullets(s, ["Completed evidence is separated from hypotheses and planned work.",
                "Every proposed mechanism must beat the strongest fair baseline across seeds.",
                "Validation guides design; the historically reused test split is not treated as untouched confirmation."], y=4.05, size=17)

    s = prs.slides.add_slide(blank); slide_title(s, "Fixed data and evaluation boundary")
    ppt_table(s, [["Role", "Images", "Patients", "Use"], ["BRSET train", "11,372", "5,966", "Source fitting"],
                  ["mBRSET train", "3,402", "899", "Target fitting"], ["mBRSET validation", "725", "193", "Checkpoint + thresholds"],
                  ["mBRSET test", "732", "193", "Gated assessment; reused historically"]], .65, 1.55, 12.0, 2.5, [2.7, 1.2, 1.2, 4.9], 14)
    ppt_box(s, .85, 4.55, 3.55, 1.25, "B0\nmBRSET only", PALE, BLUE, 18, True)
    ppt_box(s, 4.9, 4.55, 3.55, 1.25, "B1\nBRSET + mBRSET jointly", LIGHT_GREEN, GREEN, 18, True)
    ppt_box(s, 8.95, 4.55, 3.55, 1.25, "B2\nBRSET then mBRSET", PALE, BLUE, 18, True)
    textbox(s, .9, 6.25, 11.5, .45, "All arms: 5,775 updates, effective batch 64, same model/loss/augmentation/selection rule", 15, True, MUTED, PP_ALIGN.CENTER)

    s = prs.slides.add_slide(blank); slide_title(s, "Step 0: fitted degradation matches selected appearance statistics")
    ppt_table(s, [["Preprocessing", "Unchanged", "Generic", "Fitted"], ["Legacy 512", "6.742", "10.092", "1.261"],
                  ["560 resize / 512 crop", "8.474", "12.011", "4.294"]], 1.15, 1.7, 11.0, 1.65, [3.4, 2, 2, 2], 16)
    bullets(s, ["Lower distance means closer agreement with five measured appearance statistics.",
                "Frozen parameters improved the distance on patient-disjoint training-only images.",
                "This does not establish lesion preservation, visual realism, or classifier benefit."], y=4.05, size=18)

    s = prs.slides.add_slide(blank); slide_title(s, "Step 2: full-pool baseline results across three seeds", "Mean ± sample SD on the common 732-image mBRSET assessment split")
    means = summary["aggregate"]
    table_rows = [["Arm", "DR F1", "DR AUROC", "ME F1", "ME AUROC"]]
    for arm in ("B0", "B1", "B2"):
        vals=[]
        for label, metric in [("diabetic_retinopathy","f1_positive"),("diabetic_retinopathy","auroc"),("macular_edema","f1_positive"),("macular_edema","auroc")]:
            x=means[arm][label][metric]; vals.append(f"{x['mean']:.4f} ± {x['sample_sd']:.4f}")
        table_rows.append([arm]+vals)
    ppt_table(s, table_rows, .55, 1.55, 12.2, 2.15, [1.2, 2.4, 2.4, 2.4, 2.4], 14)
    ppt_box(s, .85, 4.25, 5.65, 1.2, "DR: B1 − B0 = +0.0233\npositive in 3/3 seeds", LIGHT_GREEN, GREEN, 20, True)
    ppt_box(s, 6.85, 4.25, 5.65, 1.2, "ME: B1 − B0 = +0.0379\npositive in 3/3 seeds", LIGHT_GREEN, GREEN, 20, True)
    textbox(s, .85, 5.95, 11.65, .55, "B1 is the Step-3 reference. This is an engineering result, not a novelty or clinical claim.", 17, True, AMBER, PP_ALIGN.CENTER)

    s = prs.slides.add_slide(blank); slide_title(s, "What Step 2 established—and what it did not")
    ppt_box(s, .7, 1.55, 5.8, .55, "Established", LIGHT_GREEN, GREEN, 20, True)
    bullets(s, ["Nine full training runs completed.", "Independent metric, cohort, hash and exposure checks passed.",
                "Joint training is the strongest ordinary baseline."], x=.85, y=2.3, w=5.45, h=3.6, size=17)
    ppt_box(s, 6.85, 1.55, 5.8, .55, "Still unresolved", LIGHT_AMBER, AMBER, 20, True)
    bullets(s, ["Why joint training helps.", "Clinical significance and operating point.",
                "Untouched external generalization.", "A defensible novel mechanism."], x=7.0, y=2.3, w=5.45, h=3.6, size=17)

    s = prs.slides.add_slide(blank); slide_title(s, "Step 3: isolate domain frequency and label composition")
    if len(control_rows) > 1:
        ref_dr, ref_me = float(control_rows[0][1]), float(control_rows[0][3])
        result_rows = [["Arm", "DR F1", "Δ", "DR AUROC", "ME F1", "Δ", "ME AUROC"]]
        for row in control_rows:
            dr, me = float(row[1]), float(row[3])
            result_rows.append([row[0], row[1], "—" if row[0].startswith("B1") else f"{dr-ref_dr:+.4f}", row[2],
                                row[3], "—" if row[0].startswith("B1") else f"{me-ref_me:+.4f}", row[4]])
        ppt_table(s, result_rows, .45, 1.48, 12.45, 2.55, [1.4, 1.25, .9, 1.5, 1.25, .9, 1.5], 12)
    else:
        ppt_table(s, [["", "Natural label mix", "Target-matched DR/ME mix"], ["Natural domain ratio", "B1 existing", "C2"],
                      ["50% BRSET / 50% mBRSET", "C1", "C3"]], 1.15, 1.55, 11.0, 2.25, [3.4, 3.8, 3.8], 16)
    status = "PASSED" if preflight and preflight.get("pass") else "PENDING"
    if STEP3_C1_CROSS.exists():
        cross = json.loads(STEP3_C1_CROSS.read_text())
        dr = cross["aggregate"]["diabetic_retinopathy"]["f1_positive"]
        me = cross["aggregate"]["macular_edema"]["f1_positive"]
        ppt_box(s, 1.0, 4.35, 3.45, 1.0, f"C1−B1 DR F1\n{dr['difference_mean']:+.4f} mean", PALE, BLUE, 18, True)
        ppt_box(s, 4.95, 4.35, 3.45, 1.0, f"C1−B1 ME F1\n{me['difference_mean']:+.4f} mean", PALE, BLUE, 18, True)
        decision_text = "B1 retained\nStep 3 complete" if STEP3_FINAL.exists() else "Three seeds\nvalidation only"
        ppt_box(s, 8.9, 4.35, 3.45, 1.0, decision_text, LIGHT_GREEN, GREEN, 18, True)
    else:
        ppt_box(s, 1.0, 4.35, 3.45, 1.0, "Only the sampler changes", PALE, BLUE, 18, True)
        ppt_box(s, 4.95, 4.35, 3.45, 1.0, "Validation-only screening", PALE, BLUE, 18, True)
        ppt_box(s, 8.9, 4.35, 3.45, 1.0, f"Allocated preflight\n{status}", LIGHT_GREEN if status=="PASSED" else LIGHT_AMBER, GREEN if status=="PASSED" else AMBER, 18, True)
    textbox(s, .85, 5.80, 11.7, .75, run_state.capitalize() + ".", 16, True, INK, PP_ALIGN.CENTER)

    s = prs.slides.add_slide(blank); slide_title(s, "Step 4: target-calibrated augmentation screen")
    audit_bullet = "Training-only audit: label standardization changed every appearance gap by <0.05 target SD."
    if STEP4_REFIT.exists():
        refit = json.loads(STEP4_REFIT.read_text())
        fit_bullet = (f"Held-out appearance distance: unchanged {refit['identity_validation']['distance_mean']:.3f}; "
                      f"full fit {refit['arms']['full']['validation']['distance_mean']:.3f}; "
                      f"no overlays {refit['arms']['overlay_free']['validation']['distance_mean']:.3f}.")
    else:
        fit_bullet = "Training-only transform refit is pending."
    if STEP4_SEED0.exists():
        screen = json.loads(STEP4_SEED0.read_text())
        ref = screen["reference"]["metrics"]
        result_rows = [["Arm", "DR F1", "Δ", "DR AUC", "ME F1", "Δ", "ME AUC", "Gate"]]
        result_rows.append(["B1", f"{ref['diabetic_retinopathy']['f1_positive']:.4f}", "—", f"{ref['diabetic_retinopathy']['auroc']:.4f}", f"{ref['macular_edema']['f1_positive']:.4f}", "—", f"{ref['macular_edema']['auroc']:.4f}", "Ref"])
        for arm, record in screen["arms"].items():
            metric, delta = record["metrics"], record["delta_from_B1"]
            result_rows.append([arm, f"{metric['diabetic_retinopathy']['f1_positive']:.4f}", f"{delta['diabetic_retinopathy']['f1_positive']:+.4f}", f"{metric['diabetic_retinopathy']['auroc']:.4f}", f"{metric['macular_edema']['f1_positive']:.4f}", f"{delta['macular_edema']['f1_positive']:+.4f}", f"{metric['macular_edema']['auroc']:.4f}", "PASS" if record["replication_screen_pass"] else "FAIL"])
        ppt_table(s, result_rows, .45, 1.35, 12.45, 2.55, [1.05, 1.25, .85, 1.3, 1.25, .85, 1.3, 1.15], 11)
        bullets(s, [audit_bullet, fit_bullet, "All three DR F1 changes were positive but below +0.01; ME was flat or slightly lower.", "No arm passed the frozen replication gate; no Step-4 test assessment was performed."], y=4.15, h=1.45, size=14)
        ppt_box(s, 1.2, 5.85, 10.9, .72, "Conclusion: global appearance matching did not produce a large diagnostic gain", LIGHT_AMBER, AMBER, 18, True)
    else:
        bullets(s, ["Reference fixed: natural joint B1 after Step-3 cross-seed validation controls.", audit_bullet, fit_bullet,
                    "Next: test the three frozen augmentation arms.", "Require consistent diagnostic gain and preservation evidence."], y=1.35, h=3.95, size=16)
        ppt_box(s, 1.2, 5.65, 10.9, .8, "Current novelty status: candidate hypothesis, not demonstrated contribution", LIGHT_AMBER, AMBER, 19, True)

    s = prs.slides.add_slide(blank); slide_title(s, "Dong requests and world-model decision")
    ppt_box(s, .65, 1.45, 3.8, .6, "Delivered", LIGHT_GREEN, GREEN, 19, True)
    bullets(s, ["Matched baseline protocol", "DR/ME prevalence and sampling controls", "Source-verified degradation examples"], x=.75, y=2.25, w=3.55, h=3.1, size=16)
    ppt_box(s, 4.78, 1.45, 3.8, .6, "Open", LIGHT_AMBER, AMBER, 19, True)
    bullets(s, ["A mechanism that beats B1 repeatedly", "Diagnostic-preservation evidence", "Fair latent-transition or routing study"], x=4.88, y=2.25, w=3.55, h=3.1, size=16)
    ppt_box(s, 8.91, 1.45, 3.8, .6, "World-model gate", PALE, BLUE, 19, True)
    bullets(s, ["CheXWorld, CVPR 2025: closest conceptual match", "GenDeg, CVPR 2025: separate diffusion option", "Freeze a bounded content-aware test before training"], x=9.01, y=2.25, w=3.55, h=3.1, size=16)
    textbox(s, .85, 6.15, 11.65, .45, "Internal paper targets: complete draft 15 October • hard freeze 19 October • official deadline 26 October", 16, True, BLUE, PP_ALIGN.CENTER)

    prs.save(PPTX)


def verify(summary_inputs):
    with zipfile.ZipFile(DOCX) as archive:
        doc_xml = "".join(archive.read(name).decode("utf-8", "ignore") for name in archive.namelist() if name.endswith(".xml"))
    with zipfile.ZipFile(PPTX) as archive:
        ppt_xml = "".join(archive.read(name).decode("utf-8", "ignore") for name in archive.namelist() if name.startswith("ppt/slides/") and name.endswith(".xml"))
    doc = Document(DOCX); prs = Presentation(PPTX)
    bounds_pass = all(
        shape.left >= 0 and shape.top >= 0
        and shape.left + shape.width <= prs.slide_width
        and shape.top + shape.height <= prs.slide_height
        for slide in prs.slides for shape in slide.shapes
    )
    result = {"pass": True, "generated": UPDATED, "docx": {"path": str(DOCX), "sha256": sha(DOCX),
              "sections": len(doc.sections), "tables": len(doc.tables), "paragraphs": len(doc.paragraphs),
              "times_new_roman_present": FONT in doc_xml},
              "pptx": {"path": str(PPTX), "sha256": sha(PPTX), "slides": len(prs.slides),
              "times_new_roman_present": FONT in ppt_xml, "all_shape_bounds_inside_slide": bounds_pass}, "input_sha256": summary_inputs,
              "privacy": "No dataset images, patient identifiers, private manifests, or per-image predictions embedded."}
    if len(prs.slides) != 9 or len(doc.tables) < 7 or FONT not in doc_xml or FONT not in ppt_xml or not bounds_pass:
        result["pass"] = False
    CHECKS.write_text(json.dumps(result, indent=2) + "\n")
    if not result["pass"]: raise RuntimeError(result)
    return result


def main():
    summary = json.loads(RESULTS.read_text()); design = json.loads(DESIGN.read_text())
    preflight = json.loads(PREFLIGHT.read_text()) if PREFLIGHT.exists() else None
    build_docx(summary, design, preflight); build_pptx(summary, preflight)
    inputs = {str(path.relative_to(ROOT)): sha(path) for path in (RESULTS, SEED2, DESIGN) if path.exists()}
    if PREFLIGHT.exists(): inputs[str(PREFLIGHT.relative_to(ROOT))] = sha(PREFLIGHT)
    if STEP3_SEED0.exists(): inputs[str(STEP3_SEED0.relative_to(ROOT))] = sha(STEP3_SEED0)
    if STEP3_C1_CROSS.exists(): inputs[str(STEP3_C1_CROSS.relative_to(ROOT))] = sha(STEP3_C1_CROSS)
    if STEP3_FINAL.exists(): inputs[str(STEP3_FINAL.relative_to(ROOT))] = sha(STEP3_FINAL)
    if STEP4_AUDIT.exists(): inputs[str(STEP4_AUDIT.relative_to(ROOT))] = sha(STEP4_AUDIT)
    if STEP4_REFIT.exists(): inputs[str(STEP4_REFIT.relative_to(ROOT))] = sha(STEP4_REFIT)
    if STEP4_PREFLIGHT.exists(): inputs[str(STEP4_PREFLIGHT.relative_to(ROOT))] = sha(STEP4_PREFLIGHT)
    if STEP4_SEED0.exists(): inputs[str(STEP4_SEED0.relative_to(ROOT))] = sha(STEP4_SEED0)
    print(json.dumps(verify(inputs), indent=2))


if __name__ == "__main__":
    main()
