"""Build advisor-only DOCX/PPTX companions containing licensed fundus images.

The output directory is ignored by Git. Image decoding/cropping is restricted
to an allocated Slurm compute node under the project resource policy.
"""
from __future__ import annotations

import hashlib
import json
import os
import socket
import zipfile
from pathlib import Path

import numpy as np
from PIL import Image
from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches as PInches, Pt as PPt

ROOT = Path("/home/users/sthummala2/brset-convnextv2")
REPORTS = ROOT / "research/codex/reports"
STEP4 = ROOT / "research/codex/step4"
PRIVATE = REPORTS / "private"
GALLERY = STEP4 / "private/step4_transform_gallery.png"
PREFLIGHT = STEP4 / "preflight.json"
PUBLIC_DOCX = REPORTS / "BRSET_mBRSET_Research_Record.docx"
PUBLIC_PPTX = REPORTS / "BRSET_mBRSET_Research_Update.pptx"
OUT_DOCX = PRIVATE / "BRSET_mBRSET_Advisor_Record_with_Images.docx"
OUT_PPTX = PRIVATE / "BRSET_mBRSET_Advisor_Update_with_Images.pptx"
CHECKS = PRIVATE / "advisor_visual_checks.json"
FONT = "Times New Roman"
LABELS = ["BRSET original", "FundusAug component", "Target-fitted full",
          "Target-fitted no overlays", "Real mBRSET (unpaired)"]
# Gallery row 1 is authentic, but its raw BRSET file contains a prominent
# horizontal acquisition seam. Keep it in the private audit and omit it from
# the advisor-facing layout because it distracts from the transform comparison.
DISPLAY_ROW_INDICES = (2, 3, 4)


def sha(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def node_check() -> None:
    if not os.environ.get("SLURM_JOB_ID") or "login" in socket.gethostname().lower():
        raise RuntimeError("Private image report must run on an allocated Slurm node")


def set_doc_font(run, size=10, bold=False, italic=False) -> None:
    run.font.name = FONT
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), FONT)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic


def add_doc_text(paragraph, text, size=10, bold=False, italic=False):
    run = paragraph.add_run(text)
    set_doc_font(run, size, bold, italic)
    return run


def row_crops() -> list[Path]:
    with Image.open(GALLERY) as raw:
        image = raw.convert("RGB")
        array = np.asarray(image)
    dark_fraction = np.any(array < 245, axis=2).mean(axis=1)
    active = np.flatnonzero(dark_fraction > 0.25)
    groups = []
    start = previous = int(active[0])
    for value in active[1:]:
        value = int(value)
        if value != previous + 1:
            if previous - start + 1 > 100:
                groups.append((start, previous + 1))
            start = value
        previous = value
    if previous - start + 1 > 100:
        groups.append((start, previous + 1))
    if len(groups) != 4:
        raise ValueError(f"Expected four gallery image rows, found {groups}")
    outputs = []
    for index, (top, bottom) in enumerate(groups, start=1):
        if index not in DISPLAY_ROW_INDICES:
            continue
        path = PRIVATE / f"step4_gallery_row_{index}.png"
        image.crop((0, max(0, top - 2), image.width, min(image.height, bottom + 2))).save(path)
        outputs.append(path)
    return outputs


def add_landscape_section(document: Document, rows: list[Path], start_index: int) -> None:
    section = document.add_section(WD_SECTION.NEW_PAGE)
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = section.page_height, section.page_width
    section.top_margin = section.bottom_margin = Inches(0.45)
    section.left_margin = section.right_margin = Inches(0.5)
    heading = document.add_paragraph()
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_doc_text(heading, "Private Advisor Appendix: Step-4 Transformation Examples", 12, True)
    caption = document.add_paragraph()
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_doc_text(caption, "Each row uses one BRSET training image for columns 1–4. Column 5 is an unrelated mBRSET training image and is not paired ground truth.", 8.5, italic=True)
    table = document.add_table(rows=1, cols=5)
    table.style = "Table Grid"
    for cell, label in zip(table.rows[0].cells, LABELS):
        cell.text = ""
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_doc_text(paragraph, label, 7.5, True)
    for offset, image_path in enumerate(rows):
        p = document.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(2)
        p.add_run().add_picture(str(image_path), width=Inches(9.65))
        note = document.add_paragraph()
        note.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_doc_text(note, f"Example {start_index + offset}: qualitative implementation check only; diagnostic preservation is not established.", 8, italic=True)


def build_docx(rows: list[Path]) -> None:
    document = Document(PUBLIC_DOCX)
    add_landscape_section(document, rows[:2], 1)
    add_landscape_section(document, rows[2:], 3)
    document.save(OUT_DOCX)


def ppt_text(slide, x, y, w, h, value, size=15, bold=False, color=RGBColor(30, 30, 30), align=PP_ALIGN.CENTER):
    shape = slide.shapes.add_textbox(PInches(x), PInches(y), PInches(w), PInches(h))
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    run = paragraph.add_run()
    run.text = value
    run.font.name = FONT
    run.font.size = PPt(size)
    run.font.bold = bold
    run.font.color.rgb = color


def add_visual_slide(prs: Presentation, row_paths: list[Path], start_index: int) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    ppt_text(slide, .55, .23, 12.2, .55, "Step-4 original and transformed training examples", 24, True, RGBColor(31, 78, 121))
    label_width = 12.15 / 5
    for block, path in enumerate(row_paths):
        label_y = 1.03 if block == 0 else 4.10
        image_y = 1.37 if block == 0 else 4.44
        for column, label in enumerate(LABELS):
            ppt_text(slide, .59 + column * label_width, label_y, label_width, .3, label, 10, True)
        slide.shapes.add_picture(str(path), PInches(.59), PInches(image_y), width=PInches(12.15), height=PInches(2.24))
        ppt_text(slide, .8, image_y + 2.27, 11.7, .25,
                 f"Example {start_index + block}: mBRSET reference is unpaired; visual similarity does not prove lesion preservation.",
                 9, False, RGBColor(95, 95, 95))


def build_pptx(rows: list[Path]) -> None:
    presentation = Presentation(PUBLIC_PPTX)
    add_visual_slide(presentation, rows[:2], 1)
    add_visual_slide(presentation, rows[2:], 3)
    presentation.save(OUT_PPTX)


def verify(rows: list[Path]) -> None:
    preflight = json.loads(PREFLIGHT.read_text())
    if sha(GALLERY) != preflight["gallery_sha256"]:
        raise ValueError("Gallery hash does not match allocated preflight")
    document = Document(OUT_DOCX)
    presentation = Presentation(OUT_PPTX)
    bounds = all(shape.left >= 0 and shape.top >= 0 and
                 shape.left + shape.width <= presentation.slide_width and
                 shape.top + shape.height <= presentation.slide_height
                 for slide in presentation.slides for shape in slide.shapes)
    with zipfile.ZipFile(OUT_DOCX) as archive:
        doc_media = [name for name in archive.namelist() if name.startswith("word/media/")]
    with zipfile.ZipFile(OUT_PPTX) as archive:
        ppt_media = [name for name in archive.namelist() if name.startswith("ppt/media/")]
    result = {
        "pass": len(rows) == 3 and len(document.sections) >= 3 and len(presentation.slides) == 11 and bounds,
        "scope": "private advisor companion; contains licensed training images; do not commit or distribute publicly",
        "source_gallery_sha256": sha(GALLERY),
        "row_crop_sha256": {path.name: sha(path) for path in rows},
        "docx": {"path": str(OUT_DOCX), "sha256": sha(OUT_DOCX), "sections": len(document.sections), "media_files": len(doc_media)},
        "pptx": {"path": str(OUT_PPTX), "sha256": sha(OUT_PPTX), "slides": len(presentation.slides), "media_files": len(ppt_media), "all_shape_bounds_inside_slide": bounds},
        "display_policy": "Three clean rows shown; authentic gallery row 1 omitted because the raw BRSET file contains a prominent horizontal acquisition seam.",
        "limitations": "Examples verify implementation and illustrate appearance changes; they do not establish paired realism or diagnostic preservation.",
    }
    CHECKS.write_text(json.dumps(result, indent=2, allow_nan=False) + "\n")
    if not result["pass"]:
        raise RuntimeError(result)
    print(json.dumps(result, indent=2))


def main() -> None:
    node_check()
    PRIVATE.mkdir(parents=True, exist_ok=True)
    rows = row_crops()
    build_docx(rows)
    build_pptx(rows)
    verify(rows)


if __name__ == "__main__":
    main()
