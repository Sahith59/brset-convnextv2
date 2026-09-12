"""Three-page Times New Roman advisor report, with reproducible figure provenance."""
from pathlib import Path
import sys,json,hashlib,importlib.util,io,shutil
sys.path.insert(0,'/home/users/sthummala2/.cache/brset-report-tools/python')
import numpy as np,pandas as pd
from PIL import Image
from reportlab.pdfgen import canvas
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.platypus import SimpleDocTemplate,Paragraph,Spacer,Table,TableStyle,PageBreak,Image as RLImage,KeepTogether,Flowable
from reportlab.lib.pagesizes import A4
from docx import Document
from docx.shared import Pt, Inches
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import pymupdf as fitz
R=Path(__file__).resolve().parents[3]; O=R/'research/codex/updates/2026-09-10'; V=O/'revision_2'; V.mkdir(exist_ok=True)
F=Path('/home/users/sthummala2/.cache/brset-report-tools/fonts')
for name,file in [('TNR','Times.TTF'),('TNR-Bold','Timesbd.TTF'),('TNR-Italic','Timesi.TTF'),('TNR-BoldItalic','Timesbi.TTF')]: pdfmetrics.registerFont(TTFont(name,str(F/file)))
pdfmetrics.registerFontFamily('TNR',normal='TNR',bold='TNR-Bold',italic='TNR-Italic',boldItalic='TNR-BoldItalic')
styles={
 'body':ParagraphStyle('Body',fontName='TNR',fontSize=10.5,leading=13,spaceAfter=6),
 'title':ParagraphStyle('Title',fontName='TNR-Bold',fontSize=17,leading=20,spaceAfter=10),
 'h1':ParagraphStyle('H1',fontName='TNR-Bold',fontSize=12.5,leading=15,spaceBefore=6,spaceAfter=6),
 'h2':ParagraphStyle('H2',fontName='TNR-BoldItalic',fontSize=10.7,leading=13,spaceBefore=5,spaceAfter=4),
 'small':ParagraphStyle('Small',fontName='TNR',fontSize=9.2,leading=11,spaceAfter=5),
 'cell':ParagraphStyle('Cell',fontName='TNR',fontSize=9.5,leading=11.2),
 'center':ParagraphStyle('Center',fontName='TNR',fontSize=11,leading=16,alignment=TA_CENTER,spaceAfter=7)}
def P(text,style='body'): return Paragraph(text,styles[style])
def table(rows,widths):
 data=[[P(str(x),'cell') for x in row] for row in rows]; t=Table(data,colWidths=widths,hAlign='LEFT',repeatRows=1)
 t.setStyle(TableStyle([('FONTNAME',(0,0),(-1,-1),'TNR'),('VALIGN',(0,0),(-1,-1),'TOP'),('BACKGROUND',(0,0),(-1,0),colors.HexColor('#ededed')),('LINEBELOW',(0,0),(-1,0),.6,colors.grey),('LINEBELOW',(0,1),(-1,-1),.3,colors.HexColor('#d0d0d0')),('LEFTPADDING',(0,0),(-1,-1),5),('RIGHTPADDING',(0,0),(-1,-1),5),('TOPPADDING',(0,0),(-1,-1),5),('BOTTOMPADDING',(0,0),(-1,-1),5)])); return t
spec=importlib.util.spec_from_file_location('ops',R/'scripts/55_degradation_ops.py'); ops=importlib.util.module_from_spec(spec); spec.loader.exec_module(ops)
params=json.loads((R/'results/fitted_degradation_params.json').read_text()); manifest=pd.read_csv(O/'image_manifest_private.csv'); panels=[]; records=[]
paths={'BRSET':R/'data/finetune_multilabel/train','mBRSET':R/'data/finetune_mbrset_multilabel/train'}
# Validate identities against source manifests, resolved dataset roots, and raw label tables.
rawpaths={'BRSET':Path('/data/users4/nshaik3/Datasets/BRSET/physionet.org/files/brazilian-ophthalmological/1.0.1/labels_brset.csv'),'mBRSET':Path('/data/users4/nshaik3/Datasets/mBRSET/physionet.org/files/mbrset/1.0/labels_mbrset.csv')}
rawsets={}
for dom,p in rawpaths.items():
 raw=pd.read_csv(p)
 col='image_id' if dom=='BRSET' else 'file'
 rawsets[dom]={str(x).removesuffix('.jpg') for x in raw[col]}
for i in range(3):
 originals=[]
 for dom,split in [('BRSET','splits_multilabel.csv'),('mBRSET','splits_mbrset_multilabel.csv')]:
  row=manifest[manifest.domain==dom].iloc[i]; path=paths[dom]/row.file; resolved=path.resolve(strict=True)
  assert ('/Datasets/'+dom+'/') in str(resolved)
  assert str(row.file).removesuffix('.jpg') in rawsets[dom]
  m=pd.read_csv(R/'results'/split); rec=m[m.file==row.file]; assert len(rec)==1 and rec.iloc[0]['split']=='train'
  pid='patient_id' if dom=='BRSET' else 'patient'; assert str(rec.iloc[0][pid])==str(row.patient)
  with Image.open(resolved) as im: native=im.size; im=im.convert('RGB').resize((560,560),Image.Resampling.BILINEAR).crop((24,24,536,536))
  originals.append(im)
  records.append({'row':i+1,'domain':dom,'file':row.file,'resolved_path':str(resolved),'native_size':native,'sha256':hashlib.sha256(resolved.read_bytes()).hexdigest(),'raw_labels_membership':True,'training_manifest_membership':True})
 src,target=originals
 ims=[src,ops.degrade(src,ops.GDRNET_GENERIC,np.random.default_rng(100+i)),ops.degrade(src,params,np.random.default_rng(200+i)),target]
 panelrow=[]
 for j,im in enumerate(ims):
  p=V/f'row{i+1}_column{j+1}.png'; im.save(p); panelrow.append(p)
 panels.append(panelrow)
# Reproduce the parameter draw sequence without rerunning/refitting images.
rng=np.random.default_rng(0); ranges={'blur_sigma':(0,1.2),'light_strength':(0,.6),'n_spot':(0,4),'n_hole':(0,3),'halo':(0,.2),'brightness_gain':(.7,1.15),'contrast_gain':(.9,2),'saturation_gain':(.5,1.2),'noise_sigma':(0,.03)}
matched=[]
for t in range(150):
 p={k:float(rng.integers(lo,hi+1) if k.startswith('n_') else rng.uniform(lo,hi)) for k,(lo,hi) in ranges.items()}
 if all(np.isclose(p[k],params[k],rtol=0,atol=1e-14) for k in p): matched.append(t)
assert len(matched)==1
# Independently reconstruct each distance from saved statistic summaries.
a=pd.read_csv(O/'appearance_statistics.csv'); d=pd.read_csv(O/'appearance_distances.csv')
for row in d.itertuples():
 st=a[(a.preprocessing==row.preprocessing)&(a.arm==row.arm)&(a.degradation_seed==row.degradation_seed)].set_index('statistic')
 target=a[(a.preprocessing==row.preprocessing)&(a.arm=='mBRSET_reference')].set_index('statistic')
 calc=float((((st['mean']-target['mean'])/target['std'].clip(lower=1e-6))**2).sum()); assert abs(calc-row.distance)<1e-10
D=d.groupby(['preprocessing','arm']).distance.mean(); W=A4[0]-88
s=[]; content=[]
class DistanceEquation(Flowable):
 def __init__(self): Flowable.__init__(self); self.width=W; self.height=54
 def draw(self):
  c=self.canv; x=W/2-113
  def txt(dx,y,t,font='TNR',size=11): c.setFont(font,size); c.drawString(x+dx,y,t)
  txt(0,22,'D','TNR-Italic'); txt(8,18,'r',size=8); txt(14,22,'(A, T) =','TNR-Italic')
  txt(62,18,'∑',size=21); txt(68,39,'5',size=8); txt(63,7,'k = 1',size=8)
  txt(84,13,'[',size=31)
  txt(111,34,'μ','TNR-Italic'); txt(118,30,'A,r,k',size=8); txt(143,34,'−'); txt(159,34,'μ','TNR-Italic'); txt(166,30,'T,k',size=8)
  c.setLineWidth(.6); c.line(x+102,28,x+200,28)
  txt(108,13,'max('); txt(131,13,'s','TNR-Italic'); txt(136,9,'T,k',size=8); txt(150,13,', 10'); txt(171,19,'−6',size=8); txt(182,13,')')
  txt(205,13,']',size=31); txt(216,38,'2',size=9)
def add(text,style='body'):
 s.append(DistanceEquation() if style=='center' and text.startswith('<i>D</i>') else P(text,style)); content.append((style,text))
add('Degradation validation for BRSET–mBRSET','title')
add('<b>Result.</b> Applying the existing fitted transformation to a new training-only sample reduced the difference in five appearance summaries. <i>This is evidence of appearance matching, not yet evidence of better diabetic-retinopathy (DR) or macular-edema (ME) prediction.</i>')
add('1. What was tested','h1')
add('The source is BRSET; the target is mBRSET. The new check uses <b>64 BRSET images from 64 patients</b> and <b>64 mBRSET images from 59 patients</b>. Their patients are separate from the original measurement/fitting samples reconstructed from the saved scripts and random seeds. Parameters were <i>frozen</i>: kept unchanged rather than fitted again. No diagnosis labels, validation/test images or classifier predictions were used. [1,2]')
add('Image preparation: two explicitly defined settings','h2')
add('<b>Direct 512 × 512 resize</b> (called “legacy 512” in older files): resize the whole photograph to a 512-pixel square, using Pillow’s default bicubic interpolation (pixel estimation during resizing) for RGB images. This reproduces the earlier measurement procedure and stretches non-square photographs.')
add('<b>560 × 560 resize, then 512 × 512 center crop:</b> resize the whole photograph to a 560-pixel square with bilinear interpolation, then remove 24 pixels from each edge. This matches the classifier’s evaluation geometry before normalization. The transformation is applied <i>after</i> this preparation in the present check; random training augmentations are not included. [2]')
add('What “appearance distance” means','h2')
add('Five quantities are measured within a threshold-based retinal mask: high-frequency energy (“sharpness”), mean brightness, intensity variation (“contrast”), center-to-edge brightness difference (“falloff”), and saturation (color intensity). Their exact definitions are on page 2.')
add('<i>D</i><sub>r</sub>(<i>A,T</i>) = ∑<sub>k=1</sub><super>5</super> [(<i>μ</i><sub>A,r,k</sub> − <i>μ</i><sub>T,k</sub>) / max(<i>s</i><sub>T,k</sub>, 10<super>−6</super>)]<super>2</super>','center')
add('For measure <i>k</i>, <i>μ</i><sub>A,r,k</sub> is its average across transformed source images at random seed <i>r</i>; <i>μ</i><sub>T,k</sub> and <i>s</i><sub>T,k</sub> are the target-image mean and sample standard deviation. Dividing by target spread puts different units on a comparable scale. <b>Smaller is closer on these five averages.</b> It is not a percentage, diagnostic score or comparison of complete image distributions.')
add('Results on the new sample','h2')
rows=[['<b>Preparation</b>','<b>Unchanged<br/>BRSET</b>','<b>Local<br/>generic</b>','<b>Frozen<br/>fitted</b>']]
for mode,label in [('legacy_512','Direct 512 × 512'),('model_eval_560_crop512','560 × 560 → center 512 × 512')]: rows.append([label]+[f'{D[mode,arm]:.3f}' for arm in ['BRSET_unaltered','local_generic','fitted_frozen']])
s.append(table(rows,[W*.46,W*.18,W*.18,W*.18])); content.append(('table',rows))
add('Transform columns average <i>D</i><sub>r</sub> over seeds 0, 1 and 2 on the same images; unchanged BRSET is deterministic. Target means and spreads are recalculated for each row, so interpret comparisons <b>within a row</b>. These three transform draws do not provide uncertainty across patient samples. The local generic control is handcrafted, <b>not official FundusAug or full GDRNet</b>. [2,3]','small')
s.append(PageBreak()); content.append(('pagebreak',''))
add('2. How the degradation was produced','h1')
add('The code applies this ordered chain: <b>Gaussian blur → uneven illumination → bright/dark spots and halo → sensor-like noise → brightness, contrast and saturation adjustments.</b> Pixel values are clipped to the valid image range. Each configured operation is applied; there is no independent 50% application gate. This is an empirical simulator, not an identified physical model of the target camera. [1]')
add('Parameter search and selected settings','h2')
add('The original search measured 400 training images per dataset and evaluated <b>150 random parameter combinations on 25 BRSET training images</b>. It minimized the distance above to the saved mBRSET means/spreads. Source fitting images used direct 512-square resizing; the original target measurement used short-side resizing to 512 followed by a center crop. Continuous parameters were sampled uniformly; counts were sampled uniformly over inclusive integers. [1]')
settings=[('blur_sigma','Blur standard deviation (pixels)','0–1.2'),('light_strength','Illumination strength','0–0.6'),('n_spot','Number of bright spots','0–4'),('n_hole','Number of dark spots','0–3'),('halo','Halo amplitude','0–0.20'),('brightness_gain','Brightness multiplier','0.70–1.15'),('contrast_gain','Contrast multiplier','0.90–2.00'),('saturation_gain','Saturation multiplier','0.50–1.20'),('noise_sigma','Noise standard deviation (0–1 RGB)','0–0.030')]
pr=[['<b>Parameter (meaning)</b>','<b>Search range</b>','<b>Local generic</b>','<b>Selected fit</b>']]
for k,label,ran in settings:
 form=lambda x:str(int(x)) if k.startswith('n_') else f'{x:.4f}'
 pr.append([label,ran,form(ops.GDRNET_GENERIC[k]),form(params[k])])
s.append(table(pr,[W*.43,W*.18,W*.18,W*.21])); content.append(('table',pr))
add('Displayed settings are rounded; the saved parameter file retains full precision. A random-number generator (RNG) makes sampling reproducible. Search/sample seed: 0; transform RNG at trial <i>t</i>: <i>t</i>. The stored parameter vector matches trial '+str(matched[0])+' (zero-based) in the reconstructed draw sequence. New-sample selection seed: 20260910. New quantitative transform seeds: 0–2; figure generic/fitted seeds: 100–102 / 200–202.','small')
add('Random placement and fixed operator details','h2')
add('Let <i>R</i> be half the shorter image side. Spot centers are uniform within 15–85% of image height/width. Bright-spot Gaussian radius: 0.015–0.055<i>R</i>; dark-spot radius: 0.015–0.045<i>R</i>; amplitude: 0.10–0.30. Halo radius: 0.82<i>R</i>, width: 0.10<i>R</i>. Illumination-center offsets are within ±22% of each side; Gaussian width is 55% of the shorter side. Its multiplicative field is 1 + strength × (Gaussian − 0.5). Noise is zero-mean Gaussian per RGB channel, applied where the pre-noise mask is positive. [1]','small')
add('Exact appearance measurements','h2')
add('RGB values are scaled to [0,1]; grayscale is their channel mean. The mask contains pixels with grayscale &gt; 0.04. “Sharpness” is the variance of the grayscale Laplacian (a local intensity-change operator) within that mask; noise can increase it. Brightness is masked mean intensity; contrast is masked population standard deviation. Falloff is mean intensity at radius &lt; 25% of the shorter side minus that at radius &gt; 39%, both masked. Saturation averages (maximum RGB − minimum RGB) / maximum RGB, with a 10<super>−6</super> denominator floor. Target spread in the distance uses sample standard deviation <i>across images</i>. [1,2]','small')
add('<b>Scope:</b> the mask excludes most black background but is not an anatomical segmentation. This check does not test whether simulated spots obscure or resemble lesions. Training and clinical validation remain necessary.','small')
s.append(PageBreak()); content.append(('pagebreak',''))
add('3. Verified image examples and next step','h1')
add('Each row contains one original BRSET image, two transformations of that same image, and an <b>unrelated real mBRSET reference</b>. The first three randomly selected images from each new-sample manifest are shown; they were not selected by diagnosis or visual quality. “Original” means no added degradation, after the stated resize/crop.','small')
heads=['<b>BRSET<br/>original</b>','<b>Local generic<br/>transform</b>','<b>Frozen fitted<br/>transform</b>','<b>Real mBRSET<br/>(unpaired)</b>']
cell=W/4-5; grid=[[P(h,'cell') for h in heads]]
for ims in panels: grid.append([RLImage(str(p),width=cell,height=cell) for p in ims])
g=Table(grid,colWidths=[W/4]*4); g.setStyle(TableStyle([('FONTNAME',(0,0),(-1,-1),'TNR'),('VALIGN',(0,0),(-1,-1),'TOP'),('LEFTPADDING',(0,0),(-1,-1),2.5),('RIGHTPADDING',(0,0),(-1,-1),2.5),('TOPPADDING',(0,0),(-1,-1),2),('BOTTOMPADDING',(0,0),(-1,-1),3)])); s.append(g); content.append(('figure',panels))
add('<i>Provenance.</i> BRSET rows: img04157.jpg, img06199.jpg, img15850.jpg; mBRSET rows: 1317.1.jpg, 252.3.jpg, 927.1.jpg. All six resolve to their respective BRSET v1.0.1 / mBRSET v1.0 dataset files and match the raw label tables and training manifests. File hashes and transform seeds are retained in the internal provenance record. The panels use classifier evaluation geometry; the added spots are simulator outputs, not new annotations. [2]','small')
add('Interpretation and planned validation','h2')
add('The frozen fit improves the measured appearance objective beyond its original fitting sample. The images still show synthetic spots and remaining differences in color, framing and illumination. <b>Diagnostic preservation, classification improvement and novelty are not established.</b> Next: obtain ophthalmic review of representative lesion-bearing pairs; lock a fair training protocol; then compare ordinary augmentation, faithful FundusAug and fitted augmentation before adding distillation.','small')
add('Evidence and reproduction','h2')
add('[1] Repository: scripts/55_degradation_ops.py, 56_measure_datasets.py, 57_fit_degradation.py; results/fitted_degradation_params.json.<br/>[2] Evaluation: research/codex/tools/degradation_update.py; appearance_statistics.csv, appearance_distances.csv, validation_summary.json; revision_2/figure_provenance.json. Original samples are reconstructed from saved code/seeds, not a contemporaneous file manifest.<br/>[3] Che et al., <i>Towards Generalizable Diabetic Retinopathy Grading in Unseen Domains</i>, MICCAI 2023. <link href="https://github.com/chehx/DGDR/blob/main/dataset/fundusaug.py" color="#000000"><u>Official FundusAug implementation</u></link>.','small')
def footer(c,doc):
 c.setFont('TNR',9); c.drawRightString(A4[0]-44,25,f'{doc.page} / 3')
class TNRCanvas(canvas.Canvas):
 def __init__(self,*a,**k): k['initialFontName']='TNR'; super().__init__(*a,**k)
path=V/'Degradation_Update.pdf'; doc=SimpleDocTemplate(str(path),pagesize=A4,rightMargin=44,leftMargin=44,topMargin=35,bottomMargin=37,title='Degradation validation for BRSET–mBRSET',author='Sahith Reddy Thummala')
doc.build(s,onFirstPage=footer,onLaterPages=footer,canvasmaker=TNRCanvas)
pdf=fitz.open(path); assert len(pdf)==3,f'Expected 3 pages, got {len(pdf)}'
fonts=[]
for i,page in enumerate(pdf):
 fonts.extend([f[3] for f in page.get_fonts()]); page.get_pixmap(matrix=fitz.Matrix(1.5,1.5)).save(V/f'page_{i+1}.png')
 for block in page.get_text('dict')['blocks']:
  if 'lines' not in block: continue
  for line in block['lines']:
   for span in line['spans']:
    assert 'TimesNewRoman' in span['font'].replace('-',''),span['font']
# Verify decoded embedded image pixels equal generated lossless panels.
embedded=[]
for xref,*_ in pdf[2].get_images(full=True):
 obj=pdf.extract_image(xref); im=Image.open(io.BytesIO(obj['image'])).convert('RGB'); embedded.append(hashlib.sha256(np.asarray(im).tobytes()).hexdigest())
expected=[hashlib.sha256(np.asarray(Image.open(p).convert('RGB')).tobytes()).hexdigest() for row in panels for p in row]
assert sorted(embedded)==sorted(expected),(len(embedded),len(expected))
prov={'source_images':records,'figure_generic_seeds':[100,101,102],'figure_fitted_seeds':[200,201,202],'preprocessing':'RGB; PIL bilinear resize560x560; crop(24,24,536,536); then transform','fitted_parameter_trial':matched[0],'all_12_embedded_panels_pixel_verified':True,'distance_reconstruction_matches':True,'pdf_pages':len(pdf),'pdf_fonts':sorted(set(fonts)),'font_source':(F/'source.txt').read_text()}
(V/'figure_provenance.json').write_text(json.dumps(prov,indent=2)); (V/'report_text.txt').write_text('\n'.join(p.get_text() for p in pdf))
# Editable companion: same material and explicit TNR styles; PDF is the verified pagination master.
d=Document(); sec=d.sections[0]; sec.page_width=Inches(8.27); sec.page_height=Inches(11.69); sec.left_margin=sec.right_margin=Inches(44/72); sec.top_margin=Inches(35/72); sec.bottom_margin=Inches(37/72)
for name,size,bold,italic in [('Normal',10.5,False,False),('Title',17,True,False),('Heading 1',12.5,True,False),('Heading 2',10.7,True,True)]:
 st=d.styles[name]; st.font.name='Times New Roman'; st.font.size=Pt(size); st.font.bold=bold; st.font.italic=italic
 st.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Times New Roman')
# Preserve bold/italic/super/sub styling in the companion.
from html.parser import HTMLParser
class Rich(HTMLParser):
 def __init__(self,p): super().__init__(); self.p=p; self.tags=[]
 def handle_starttag(self,t,attrs):
  if t=='br': self.p.add_run().add_break()
  else:self.tags.append(t)
 def handle_endtag(self,t):
  if t in self.tags:self.tags.remove(t)
 def handle_data(self,data):
  r=self.p.add_run(data); r.bold='b' in self.tags; r.italic='i' in self.tags; r.font.superscript='super' in self.tags; r.font.subscript='sub' in self.tags
for typ,x in content:
 if typ=='pagebreak': d.add_page_break()
 elif typ=='figure':
  t=d.add_table(rows=0,cols=4)
  for names in [['BRSET original','Local generic','Frozen fitted','Real mBRSET (unpaired)']]:
   for cell,txt in zip(t.add_row().cells,names): cell.text=txt
  for row in x:
   for cell,p in zip(t.add_row().cells,row): cell.paragraphs[0].add_run().add_picture(str(p),width=Inches(1.62))
 elif typ=='table':
  t=d.add_table(rows=0,cols=4); t.style='Table Grid'
  for row in x:
   for cell,txt in zip(t.add_row().cells,row): Rich(cell.paragraphs[0]).feed(str(txt))
 else:
  p=d.add_paragraph(style={'title':'Title','h1':'Heading 1','h2':'Heading 2'}.get(typ,'Normal')); Rich(p).feed(x)
  if typ=='small':
   for r in p.runs:r.font.size=Pt(9.2)
d.save(V/'Degradation_Update.docx')
print(json.dumps({'pages':len(pdf),'fonts':sorted(set(fonts)),'verified_source_images':len(records),'verified_embedded_panels':len(embedded),'matched_fit_trial':matched},indent=2))
