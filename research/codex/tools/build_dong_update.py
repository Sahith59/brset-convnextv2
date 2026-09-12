from pathlib import Path
import json,textwrap
import pandas as pd
from docx import Document
from docx.shared import Inches,Pt
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from matplotlib.backends.backend_pdf import PdfPages
O=Path(__file__).resolve().parents[1]/'updates/2026-09-10'
s=json.loads((O/'validation_summary.json').read_text()); d=pd.read_csv(O/'appearance_distances.csv').groupby(['preprocessing','arm']).distance.mean()
rows=[]
for mode,label in [('legacy_512','Legacy 512 resize'),('model_eval_560_crop512','560 resize / 512 crop')]: rows.append([label]+[f'{d[mode,arm]:.3f}' for arm in ['BRSET_unaltered','local_generic','fitted_frozen']])
intro='The frozen target-fitted transformation brings the source images closer to the target on five measured appearance statistics in a small new training-only evaluation. Diagnostic improvement has not yet been tested.'
method='The original work measured 400 training images per dataset and fitted parameters using 25 BRSET images over 150 random trials. The transformation includes blur, illumination changes, bright/dark artifacts, a halo, noise and color adjustments. The saved parameters were kept unchanged for the new check. [1]'
validation='The new check uses 64 BRSET images from 64 patients and 64 mBRSET images from 59 patients. These patients were excluded from the original measurement/fitting cohorts reconstructed from the saved scripts and seeds. Three stochastic degradation seeds were evaluated on the same images. No diagnosis labels, validation/test images, classifier inference or model training were used. [2]'
metric='Distance is the sum, over five appearance measures, of squared source-target mean differences divided by the target standard deviation squared. Lower means closer on these summaries only. Target means and scales are recomputed for each preprocessing row, so comparisons are within a row; the values are not accuracy, percentages, confidence intervals or full distribution matching.'
findings='Fitting improves this descriptive appearance objective on the new cohort under both preprocessing procedures. The fitted examples still show artificial spots and residual differences from real target images. Neither lesion preservation nor DR/ME classification benefit is established. The comparison labeled local generic is our existing handcrafted transform, not published FundusAug or full GDRNet. [2,3]'
nextp='Next, validate representative lesion-bearing pairs, align fitting with the selected classifier preprocessing, and compare ordinary augmentation, a faithful FundusAug component and fitted augmentation under one training protocol. Test augmentation before adding distillation. In parallel, control label/domain balance and training exposure as discussed in the meeting. A fair longer-training routing comparison remains a separate follow-up.'
caption='First three randomly selected BRSET examples from the new cohort. Columns 1–3 show the same source image and two transformations. Column 4 shows unrelated real mBRSET reference images: these are not paired acquisitions. Display uses 560-pixel bilinear square resize and a 512-pixel center crop. Figure transform seeds are 100–102 and 200–202; quantitative seeds are 0–2. Source files and seeds are recorded in the internal manifest. This figure supports qualitative review, not a clinical certification.'
refs=['[1] Existing project files: scripts/55_degradation_ops.py, 56_measure_datasets.py, 57_fit_degradation.py; results/fitted_degradation_params.json.','[2] New evaluation: research/codex/tools/degradation_update.py; validation_summary.json, appearance_statistics.csv and appearance_distances.csv, September 10, 2026.','[3] Che et al., Towards Generalizable Diabetic Retinopathy Grading in Unseen Domains, MICCAI 2023. Official augmentation implementation: https://github.com/chehx/DGDR/blob/main/dataset/fundusaug.py']
md=['**Degradation work update**','',intro,'',method,'',validation,'','| Preprocessing | BRSET unchanged | Local generic | Frozen fitted |','|---|---:|---:|---:|']+['| '+' | '.join(r)+' |' for r in rows]+['',metric,'',findings,'','![Degradation comparison](degradation_examples.png)','',caption,'',nextp,'','**Sources**','']+refs
(O/'TECHNICAL_NOTE.md').write_text('\n'.join(md)+'\n')
message='''Professor Dong,

Following your request to inspect and validate the degraded images, I prepared the attached comparison and a training-only check of the existing fitted transformation.

The figure shows each original BRSET image beside a local generic transformation and the target-fitted version, with unrelated real mBRSET images for reference. I kept the fitted parameters fixed and evaluated 64 new training images per dataset, excluding patients from the original measurement/fitting samples reconstructed from the scripts.

The fitted transformation reduced the five-statistic appearance distance from 6.742 to 1.261 under the original measurement resize, and from 8.474 to 4.294 under the classifier's resize-and-crop procedure. These comparisons use separately computed target statistics within each preprocessing setting. This supports appearance matching on the new sample; it does not yet demonstrate preserved diagnostic detail or improved DR/ME performance.

I also corrected the generic control's description: it is our local handcrafted transform, not a reproduction of published FundusAug. The next steps are to validate lesion preservation and compare fitted augmentation with ordinary augmentation and a faithful FundusAug control using matched training settings. I am also preparing the class-balance/training-exposure controls you raised. The longer-training routing comparison remains a separate follow-up.

I would appreciate your feedback on whether the transformed examples are plausible and whether we can arrange ophthalmic review of lesion preservation before treating them as label-preserving training pairs.

Best,
Sahith
'''
(O/'UPDATE_FOR_DONG.md').write_text('**Draft message — prepared for review; not sent**\n\n'+message)
# Editable, self-contained advisor note.
doc=Document(); sec=doc.sections[0]; sec.top_margin=sec.bottom_margin=Inches(.65); doc.styles['Normal'].font.name='Calibri'; doc.styles['Normal'].font.size=Pt(10)
doc.add_heading('Degradation work update',0)
for p in [intro,method,validation]: doc.add_paragraph(p)
t=doc.add_table(rows=1,cols=4); t.style='Light Shading Accent 1'
for c,x in zip(t.rows[0].cells,['Preprocessing','Unchanged','Local generic','Frozen fitted']): c.text=x
for row in rows:
 for c,x in zip(t.add_row().cells,row): c.text=x
for p in [metric,findings,nextp]: doc.add_paragraph(p)
doc.add_page_break(); doc.add_picture(str(O/'degradation_examples.png'),width=Inches(6.7)); doc.add_paragraph(caption)
for x in refs: doc.add_paragraph(x)
doc.save(O/'Degradation_Update.docx')
# Two-page PDF with identical evidence and a data-driven table.
with PdfPages(O/'Degradation_Update.pdf') as pdf:
 fig=plt.figure(figsize=(8.27,11.69)); fig.text(.075,.955,'Degradation work update',fontsize=18,va='top'); y=.9
 def para(txt,y):
  lines=textwrap.wrap(txt,width=103); fig.text(.075,y,'\n'.join(lines),fontsize=9.4,va='top',linespacing=1.4); return y-len(lines)*.0165-.020
 for p in [intro,method,validation]: y=para(p,y)
 ax=fig.add_axes([.075,y-.105,.85,.1]); ax.axis('off'); table=ax.table(cellText=rows,colLabels=['Preprocessing','Unchanged','Local generic','Frozen fitted'],cellLoc='center',loc='center',colWidths=[.4,.2,.2,.2]); table.auto_set_font_size(False); table.set_fontsize(9); table.scale(1,1.6)
 for (i,j),cell in table.get_celld().items():
  cell.set_edgecolor('#cccccc'); cell.set_linewidth(.5)
  if i==0: cell.set_facecolor('#eeeeee'); cell.set_text_props(weight='bold')
 y-=.14
 for p in [metric,findings,nextp]: y=para(p,y)
 assert y>.06,y
 pdf.savefig(fig); fig.savefig(O/'note_page1_preview.png',dpi=110); plt.close(fig)
 fig=plt.figure(figsize=(8.27,11.69)); ax=fig.add_axes([.035,.37,.93,.59]); ax.imshow(plt.imread(O/'degradation_examples.png')); ax.axis('off'); fig.text(.075,.335,'\n'.join(textwrap.wrap(caption,105)),fontsize=9,va='top',linespacing=1.4)
 y=.20
 for ref in refs:
  lines=textwrap.wrap(ref,105); fig.text(.075,y,'\n'.join(lines),fontsize=8.4,va='top',linespacing=1.35); y-=len(lines)*.0145+.01
 pdf.savefig(fig); fig.savefig(O/'note_page2_preview.png',dpi=110); plt.close(fig)
print('Built unsent message, note, editable DOCX and two-page PDF.')
